from __future__ import annotations

import csv
import uuid
from pathlib import Path
from typing import Dict, Any

import openpyxl
from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.report import Report
from app.models.request import Request
from app.models.enums import RequestType

REPORT_DIR = Path("/app/uploads/reports")
REPORT_DIR.mkdir(parents=True, exist_ok=True)


def _prepare_export_data(req: Request) -> Dict[str, Any]:
    """
    Prepare data for export based on request type.
    Returns dict with 'headers' and 'rows' (list of lists).
    """
    base = [
        req.id,
        req.number,
        req.title,
        req.description or "",
        req.type.value if req.type else "",
        req.status.value if req.status else "",
        req.expected_date.isoformat() if req.expected_date else "",
        req.created_by.full_name if req.created_by else "",
        req.requester.full_name if req.requester else "",
        req.executor.full_name if req.executor else "",
        req.division.name if req.division else "",
        0.0,
        "",
        req.submitted_at.isoformat() if req.submitted_at else "",
        req.closed_at.isoformat() if req.closed_at else "",
    ]

    if req.type == RequestType.PAYMENT and hasattr(req, "payment_details") and req.payment_details:
        pt = req.payment_details
        base[11] = float(pt.amount) if pt.amount else 0.0
        base[12] = pt.currency.code if pt.currency else ""
        base.extend([
            pt.payment_purpose or "",
            pt.recipient_name or "",
            pt.recipient_bank or "",
            pt.account_number_iban or "",
            pt.swift_bic or "",
            pt.payment_method.value if pt.payment_method else "",
            pt.agent.code if pt.agent else "",
        ])
    elif req.type == RequestType.PURCHASE and hasattr(req, "purchase_details") and req.purchase_details:
        pu = req.purchase_details
        base.extend([
            pu.payment_method.value if pu.payment_method else "",
            pu.buyer_company.name if pu.buyer_company else "",
            pu.markup_notes or "",
            pu.delivery_date.isoformat() if pu.delivery_date else "",
        ])
    elif req.type == RequestType.CONSULTATION and hasattr(req, "consultation_details") and req.consultation_details:
        c = req.consultation_details
        base.extend([
            c.question_description or "",
        ])

    headers = [
        "ID", "Номер", "Заголовок", "Описание", "Тип", "Статус", "Ожидаемая дата",
        "Создал", "Инициатор", "Исполнитель", "Подразделение", "Сумма", "Валюта",
        "Отправлено", "Закрыто",
        "Назначение платёжа", "Получатель", "Банк получателя", "IBAN", "SWIFT",
        "Метод оплаты", "Агент",
        "Метод оплаты (закупка)", "Покупатель", "Наценка", "Дата поставки",
        "Консультация",
    ]

    rows = [base]
    return {"headers": headers, "rows": rows}


def generate_report(db: Session, request_id: int, export_format: str, user_id: int) -> Report:
    """
    Generate a report (CSV or XLSX) for a given request and save metadata.
    Uses pure openpyxl for XLSX - no pandas dependency needed.
    """
    req = db.query(Request).filter(Request.id == request_id).first()
    if not req:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Заявка не найдена")

    data = _prepare_export_data(req)
    filename = f"{uuid.uuid4().hex}_{request_id}.{export_format}"
    out_path = REPORT_DIR / filename

    try:
        if export_format == "xlsx":
            # Create workbook and worksheet
            wb = Workbook()
            ws = wb.active
            ws.title = "Requests Report"
            
            # Write headers
            for col_idx, header in enumerate(data["headers"], 1):
                ws.cell(row=1, column=col_idx, value=header)
            
            # Write rows
            for row_idx, row_data in enumerate(data["rows"], 2):
                for col_idx, value in enumerate(row_data, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
            
            # Set column widths
            for col_idx in range(1, len(data["headers"]) + 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = max(
                    15, len(str(data["headers"][col_idx - 1])) + 2
                )
            
            wb.save(out_path)
        else:  # CSV
            with open(out_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(data["headers"])
                writer.writerows(data["rows"])
    except Exception as e:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                            detail=f"Не удалось сгенерировать отчёт: {e}")

    # Get file size
    file_size = out_path.stat().st_size

    report = Report(
        request_id=request_id,
        export_format=export_format,
        file_path=str(out_path),
        file_size_bytes=file_size,
        created_by_id=user_id,
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report


def get_report_fields() -> Dict[str, bool]:
    """
    Return dictionary of available report fields with checkboxes.
    """
    return {
        "ID": True,
        "Номер": True,
        "Заголовок": True,
        "Описание": True,
        "Тип": True,
        "Статус": True,
        "Ожидаемая дата": True,
        "Создал": True,
        "Инициатор": True,
        "Исполнитель": True,
        "Подразделение": True,
        "Сумма": True,
        "Валюта": True,
        "Отправлено": True,
        "Закрыто": True,
        "Получатель": True,
        "Банк получателя": True,
        "IBAN": True,
        "SWIFT": True,
        "Метод оплаты": True,
        "Агент": True,
        "Метод оплаты (закупка)": False,
        "Покупатель": False,
        "Наценка": False,
        "Дата поставки": False,
        "Консультация": False,
        "Способ оплаты": False,
        "Метод доставки": False,
        "Заказчик": False,
        "Исполнитель": False,
        "Дата создания": False,
        "Дата изменения": False,
        "Тема консультации": False,
        "Консультант": False,
        "Статус оплаты": False,
        "Способ оплаты": False,
        "Сумма оплаты": False,
        "Валюта оплаты": False,
        "Стоимость услуги": False,
        "Скидка": False,
        "Итого": False,
        "Дата оплаты": False,
        "Метод оплаты": False,
        "Статус оплаты": False,
        "Сумма оплаты": False,
        "Валюта оплаты": False,
        "Клиент": False,
        "Исполнитель": False,
        "Дата выполнения": False,
        "Дата закрытия": False,
        "Статус выполнения": False,
        "Причина закрытия": False,
        "Документы": False,
        "Дата регистрации документа": False,
        "Номер документа": False,
        "Тип документа": False,
        "Дата документа": False,
        "Клиентский номер": False,
        "Номер счета": False,
        "Сумма счета": False,
        "Валюта счета": False,
        "Дата выставления": False,
        "Дата оплаты": False,
        "Номер чека": False,
        "Сумма чека": False,
        "Валюта чека": False,
        "Код операции": False,
        "Описание операции": False,
        "Статус операции": False,
        "Примечание": False,
    }


def generate_report_list(db: Session, selected_fields: list, filters: dict, export_format: str, user_id: int) -> Report:
    """
    Generate a report (CSV or XLSX) for a list of requests and save metadata.
    """
    from app.models.request import Request
    from app.models.enums import RequestType
    
    # Build base query
    query = db.query(Request)
    
    # Apply filters
    if filters.get("type"):
        query = query.filter(Request.type.in_(filters["type"]))
    if filters.get("status"):
        query = query.filter(Request.status.in_(filters["status"]))
    if filters.get("division_id"):
        query = query.filter(Request.division_id.in_(filters["division_id"]))
    if filters.get("created_by_id"):
        query = query.filter(Request.created_by_id.in_(filters["created_by_id"]))
    if filters.get("requester_id"):
        query = query.filter(Request.requester_id.in_(filters["requester_id"]))
    if filters.get("executor_id"):
        query = query.filter(Request.executor_id.in_(filters["executor_id"]))
    
    # Date filtering
    date_field = filters.get("date_field", "created_at")
    if filters.get("date_from"):
        if date_field == "expected_date":
            query = query.filter(Request.expected_date >= filters["date_from"])
        elif date_field == "submitted_at":
            query = query.filter(Request.submitted_at >= filters["date_from"])
        else:
            query = query.filter(Request.created_at >= filters["date_from"])
    if filters.get("date_to"):
        if date_field == "expected_date":
            query = query.filter(Request.expected_date <= filters["date_to"])
        elif date_field == "submitted_at":
            query = query.filter(Request.submitted_at <= filters["date_to"])
        else:
            query = query.filter(Request.created_at <= filters["date_to"])
    
    # Execute query
    requests = query.all()
    
    # Prepare data
    rows = []
    for req in requests:
        row = []
        for field in selected_fields:
            if field == "ID":
                row.append(req.id)
            elif field == "Номер":
                row.append(req.number)
            elif field == "Заголовок":
                row.append(req.title)
            elif field == "Описание":
                row.append(req.description or "")
            elif field == "Тип":
                row.append(req.type.value)
            elif field == "Статус":
                row.append(req.status.value)
            elif field == "Ожидаемая дата":
                row.append(req.expected_date.isoformat() if req.expected_date else "")
            elif field == "Создал":
                row.append(req.created_by.full_name if req.created_by else "")
            elif field == "Инициатор":
                row.append(req.requester.full_name if req.requester else "")
            elif field == "Исполнитель":
                row.append(req.executor.full_name if req.executor else "")
            elif field == "Подразделение":
                row.append(req.division.name if req.division else "")
            elif field == "Сумма":
                # Для платёжных заявок — фактически исполненная сумма (BR-091),
                # если есть запись исполнения, иначе заявленная из payment_details.
                if req.type == RequestType.PAYMENT and req.payment_details:
                    ex = req.payment_details.executions[-1] if req.payment_details.executions else None
                    if ex is not None and ex.amount is not None:
                        row.append(float(ex.amount))
                    elif req.payment_details.amount is not None:
                        row.append(float(req.payment_details.amount))
                    else:
                        row.append("")
                else:
                    row.append("")
            elif field == "Валюта":
                if req.type == RequestType.PAYMENT and req.payment_details:
                    ex = req.payment_details.executions[-1] if req.payment_details.executions else None
                    if ex is not None and ex.currency:
                        row.append(ex.currency.code)
                    elif req.payment_details.currency:
                        row.append(req.payment_details.currency.code)
                    else:
                        row.append("")
                else:
                    row.append("")
            elif field == "Отправлено":
                row.append(req.submitted_at.isoformat() if req.submitted_at else "")
            elif field == "Закрыто":
                row.append(req.closed_at.isoformat() if req.closed_at else "")
            elif field == "Получатель":
                row.append(req.payment_details.recipient_name if req.type == RequestType.PAYMENT and req.payment_details else "")
            elif field == "Банк получателя":
                row.append(req.payment_details.recipient_bank if req.type == RequestType.PAYMENT and req.payment_details else "")
            elif field == "IBAN":
                row.append(req.payment_details.account_number_iban if req.type == RequestType.PAYMENT and req.payment_details else "")
            elif field == "SWIFT":
                row.append(req.payment_details.swift_bic if req.type == RequestType.PAYMENT and req.payment_details else "")
            elif field == "Метод оплаты":
                row.append(req.payment_details.payment_method.value if req.type == RequestType.PAYMENT and req.payment_details else "")
            elif field == "Агент":
                row.append(req.payment_details.agent.code if req.type == RequestType.PAYMENT and req.payment_details and req.payment_details.agent else "")
            elif field == "Метод оплаты (закупка)":
                row.append(req.purchase_details.payment_method.value if req.type == RequestType.PURCHASE and req.purchase_details else "")
            elif field == "Покупатель":
                row.append(req.purchase_details.buyer_company.name if req.type == RequestType.PURCHASE and req.purchase_details else "")
            elif field == "Наценка":
                row.append(req.purchase_details.markup_notes if req.type == RequestType.PURCHASE and req.purchase_details else "")
            elif field == "Дата поставки":
                row.append(req.purchase_details.delivery_date.isoformat() if req.type == RequestType.PURCHASE and req.purchase_details and req.purchase_details.delivery_date else "")
            elif field == "Консультация":
                row.append(req.consultation_details.question_description if req.type == RequestType.CONSULTATION and req.consultation_details else "")
            else:
                row.append("")
        rows.append(row)
    
    # Create report
    filename = f"{uuid.uuid4().hex}_{filters.get('request_id', 'all')}_{export_format}"
    out_path = REPORT_DIR / filename
    
    try:
        if export_format == "xlsx":
            wb = Workbook()
            ws = wb.active
            ws.title = "Requests Report"
            
            # Write headers
            for col_idx, header in enumerate(selected_fields, 1):
                ws.cell(row=1, column=col_idx, value=header)
            
            # Write rows
            for row_idx, row_data in enumerate(rows, 2):
                for col_idx, value in enumerate(row_data, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
            
            # Set column widths
            for col_idx in range(1, len(selected_fields) + 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = max(
                    15, len(str(selected_fields[col_idx - 1])) + 2
                )
            
            wb.save(out_path)
        else:  # CSV
            with open(out_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(selected_fields)
                writer.writerows(rows)
    except Exception as e:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                            detail=f"Не удалось сгенерировать отчёт: {e}")
    
    file_size = out_path.stat().st_size
    
    report = Report(
        request_id=None,  # Отчёт по списку — не привязан к конкретной заявке
        export_format=export_format,
        file_path=str(out_path),
        file_size_bytes=file_size,
        created_by_id=user_id,
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report
